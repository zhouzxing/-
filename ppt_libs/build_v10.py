import json, os, sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
from collections import defaultdict

OUT=sys.argv[1] if len(sys.argv)>1 else 'xingbao_village_v10.pptx'
ROOT=Path('/home/geeker/.hermes/cache/scratch/xingbao_50')
items=json.loads((ROOT/'ppt_items_v12.json').read_text())
by=defaultdict(list)
for a in items:
    p=Path(a['file'])
    if p.exists() and p.stat().st_size>30000:
        by[a['topic']].append(a)
# fallback known food image
food_fallback={'topic':'food','file':str(Path('/home/geeker/.hermes/cache/scratch/xingbao_all/chinese_food.jpg')), 'width':1024,'height':683,'kb':141,'query':'prior'}
if Path(food_fallback['file']).exists(): by['food'].append(food_fallback)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
W,H=13.333,7.5
BLACK=RGBColor(8,10,12); GOLD=RGBColor(198,166,94); CREAM=RGBColor(244,236,220); WHITE=RGBColor(255,255,255); MUTED=RGBColor(190,196,205); RED=RGBColor(174,54,54)

def slide_transition(slide):
    xml='<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="medSlow"><p:fade/></p:transition>'
    old=slide._element.find(qn('p:transition'))
    if old is not None: slide._element.remove(old)
    slide._element.insert(2, etree.fromstring(xml))

def slide():
    return prs.slides.add_slide(prs.slide_layouts[6])

def bg(s): s.background.fill.solid(); s.background.fill.fore_color.rgb=BLACK

def rect(s,l,t,w,h,color=BLACK,alpha=None,shape=MSO_SHAPE.RECTANGLE):
    sh=s.shapes.add_shape(shape,Inches(l),Inches(t),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=color
    if alpha is not None:
        sf=sh.fill._xPr.find(qn('a:solidFill')); srgb=sf.find(qn('a:srgbClr'))
        srgb.append(etree.fromstring(f'<a:alpha xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="{int(alpha*100000)}"/>'))
    sh.line.fill.background(); return sh

def tb(s,text,l,t,w,h,size=18,color=WHITE,bold=False,align=PP_ALIGN.LEFT,font='Microsoft YaHei'):
    box=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box

def gold(s,x,y,w=.9,h=.05): return rect(s,x,y,w,h,GOLD)
def footer(s,n): tb(s,'邢堡村 · 严格筛选专题版',.7,7.04,6.5,.24,7,RGBColor(140,140,140)); tb(s,str(n).zfill(2),12.3,7.04,.4,.24,7,GOLD,True,PP_ALIGN.RIGHT)
def pic(s,item,l,t,w,h):
    p=Path(item['file'])
    if p.exists() and p.stat().st_size>30000:
        return s.shapes.add_picture(str(p),Inches(l),Inches(t),Inches(w),Inches(h))
    return None
def cap(s,item,text,l,t,w):
    rect(s,l,t,w,.46,BLACK,.72); tb(s,text,l+.18,t+.12,w-.36,.22,10,GOLD,True)

def gallery(page_n, topic, n, title, captions, cols=3):
    s=slide(); bg(s); tb(s,title,.85,.6,7,.5,30,WHITE,True); gold(s,.9,1.3)
    xs=[.72,4.78,8.84] if cols==3 else [.85,7.1]
    for i,item in enumerate(by[topic][:n]):
        x=xs[i%cols]; y=1.85+(i//cols)*(5.1/cols)
        pic(s,item,x,y,3.58 if cols==3 else 5.1,1.65 if cols==3 else 2.3)
        cap(s,item,captions[i%len(captions)],x,y+(1.43 if cols==3 else 2.08),3.58 if cols==3 else 5.1)
    slide_transition(s); footer(s,page_n)

def section(n,num,title,subtitle,topic='loess'):
    s=slide(); bg(s); pic(s,by[topic][0] if by[topic] else None,0,0,W,H); rect(s,0,0,W,H,BLACK,.66)
    tb(s,num,.9,2.05,1.4,.6,38,GOLD,True); tb(s,title,.9,2.85,8.5,.8,46,WHITE,True,font='SimSun'); gold(s,.92,3.78,1.05); tb(s,subtitle,.92,4.15,9.4,.6,20,CREAM); slide_transition(s); footer(s,n)
    return s

# 1 cover
s=slide(); bg(s); pic(s,by['loess'][0],0,0,W,H); rect(s,0,0,W,H,BLACK,.45); rect(s,0,3.9,W,3.6,BLACK,.62)
tb(s,'邢堡村',.82,1.8,8,.95,60,WHITE,True,font='SimSun'); gold(s,.9,2.95,1.1); tb(s,'化峪镇 · 稷山县 · 运城',.9,3.24,7,.35,18,GOLD,True)
tb(s,'严格筛选专题版 · 去重素材库 · 标准淡入切换',.92,4.7,9,.4,20,CREAM); tb(s,'不塞错图：已剔除错配、低清、水印风险素材',.92,5.25,9,.3,13,GOLD,True); slide_transition(s); footer(s,1)
# 2 content
s=slide(); bg(s); tb(s,'目录 CONTENTS',.85,.72,5.5,.52,30,WHITE,True); gold(s,.9,1.42)
for i,it in enumerate(['01 地理概貌','02 历史溯源','03 古村落建筑','04 农耕文明','05 稷山四宝','06 古迹与非遗','07 短视频入口']):
    x=.9+(i%2)*6.2; y=2.1+(i//2)*.72; tb(s,it.split()[0],x,y,.6,.3,14,GOLD,True); tb(s,it.split()[1],x+.75,y,3.5,.3,15,CREAM,True)
slide_transition(s); footer(s,2)
section(3,'01','地理概貌','黄土塬、塬坡、沟壑：邢堡村的自然基底。')
gallery(4, 'loess', 10, '黄土塬图集', ['塬面','塬坡','沟壑','耕地'])
# history use architecture/heritage image as atmospheric
s=slide(); bg(s); pic(s,by['architecture'][0],0,0,W,H); rect(s,0,0,W,H,BLACK,.68); tb(s,'02 历史溯源',.9,.95,6.5,.55,34,WHITE,True); gold(s,.92,1.72,.95)
tb(s,'稷山因后稷教稼得名，是中华农耕文明的重要记忆现场。\n邢堡村延续塬区村落格局，保留古井、场院、老槐、民居院墙与地方记忆。\n村落不是孤立建筑，而是耕作、节庆、信仰、宗族、手艺共同沉淀的生活系统。',.92,2.35,9.1,1.8,19,CREAM); slide_transition(s); footer(s,5)
section(6,'03','古村落建筑','青砖、灰瓦、院落、砖雕：晋南民居的材料语言。','architecture')
gallery(7, 'architecture', 10, '古建筑图集', ['砖雕','青砖','屋顶','院落'])
# architecture split
s=slide(); bg(s); pic(s,by['architecture'][0],0,0,7.0,H); rect(s,0,0,7.0,H,BLACK,.12); rect(s,7.0,0,6.33,H,BLACK,.76); tb(s,'青砖古韵',7.35,2.05,4.8,.6,36,WHITE,True,font='SimSun'); gold(s,7.38,2.85,1); tb(s,'门楼、院墙、屋顶和砖雕共同构成晋南建筑的秩序与符号。',7.38,3.15,4.8,.9,18,CREAM); slide_transition(s); footer(s,8)
section(9,'04','农耕文明','塬区四季：春种、夏长、秋收、冬藏。','loess')
section(10,'05','稷山四宝','麻花、板枣与晋南风味，是邢堡村走向全国的文化名片。','banzao')
# Food page conservative: one fallback generic fried-food image + text only
s=slide(); bg(s); pic(s,by['food'][0] if by['food'] else None,0,0,W,H); rect(s,0,0,W,H,BLACK,.58); tb(s,'麻花工艺与传播记忆点',.9,.8,7,.55,32,WHITE,True); gold(s,.92,1.55,1)
for i,st in enumerate(['和面','醒面','盘条','拧股','入油','出锅']):
    x=.9+i*1.92; rect(s,x,2.75,1.35,1.35,GOLD,.18,MSO_SHAPE.OVAL); tb(s,f'0{i+1}',x+.32,2.96,.7,.24,10,GOLD,True,PP_ALIGN.CENTER); tb(s,st,x+.16,3.3,.95,.24,13,WHITE,True,PP_ALIGN.CENTER)
tb(s,'注：因网上“稷山麻花”高清无水印图难辨真伪，本版不强行使用可疑麻花图，保留工艺叙事与扫码入口。',.9,5.55,10.8,.45,15,CREAM); slide_transition(s); footer(s,11)
gallery(12,'banzao',12,'稷山板枣图集',['枣林','板枣','枣花','采摘'])
s=slide(); bg(s); pic(s,by['banzao'][0],0,0,W,H); rect(s,0,0,W,H,BLACK,.55); rect(s,.8,.8,5.3,5.35,BLACK,.72); tb(s,'稷山板枣',1.1,1.55,4,.6,36,WHITE,True,font='SimSun'); gold(s,1.12,2.35,.95); tb(s,'唐代古枣林 · 万亩枣花香 · 晒制板枣 · 农遗记忆',1.12,2.75,4.2,1.5,19,CREAM); tb(s,'可用枣花节、采摘节、晒枣场、礼盒伴手礼组成四季内容日历。',8.25,4.6,4.3,.8,18,CREAM); slide_transition(s); footer(s,13)
section(14,'06','古迹与非遗','古迹、民俗、手作技艺：从历史现场到传播素材。','heritage')
gallery(15,'heritage',15,'古迹与民俗图集',['古迹','民俗','技艺'])
gallery(16,'village',16,'古村落与生活场景',['院落','村庄','塬村'])
# QR page
QR=Path('/home/geeker/.hermes/cache/scratch/xingbao_v3_imgs')
s=slide(); bg(s); pic(s,by['village'][0],0,0,W,H); rect(s,0,0,W,H,BLACK,.72); tb(s,'07 短视频与扫码传播',.85,.72,8.5,.55,28,CREAM,True); gold(s,.9,1.43,1)
x=1.15
for label,file,url in [('抖音','qr_douyin.png','douyin.com/search/稷山'),('视频号','qr_videochannel.png','videos.wechat.qq.com'),('小红书','qr_xhs.png','xiaohongshu.com/search/稷山'),('麻花','qr_mahua.png','1588.tv/techan/3545'),('稷王庙','qr_jiwang.png','weibo.com/山西文旅')]:
    qf=QR/file
    if qf.exists():
        s.shapes.add_picture(str(qf),Inches(x),Inches(2.3),Inches(1.55),Inches(1.55)); tb(s,label,x,4.02,1.55,.26,10,GOLD,True,PP_ALIGN.CENTER); tb(s,url,x-.18,4.32,1.9,.55,6,MUTED,False,PP_ALIGN.CENTER)
    x+=2.25
tb(s,'扫码查看抖音/视频号/小红书相关内容',.95,5.55,8,.35,16,CREAM); slide_transition(s); footer(s,17)
# video ideas
s=slide(); bg(s); tb(s,'短视频选题库',.85,.62,6,.5,30,WHITE,True); gold(s,.9,1.32)
for i,(topic,cap) in enumerate([('loess','塬上麦浪 15秒'),('banzao','板枣采摘 vlog'),('food','麻花出锅慢镜头'),('architecture','古院光影走位'),('village','一日邢堡生活'),('heritage','非遗民俗片段')]):
    y=1.85+i*.78; pic(s,by[topic][0] if by[topic] else None,.85,y,1.1,.62); rect(s,2.15,y,9.9,.62,CREAM,.06); tb(s,f'0{i+1}',2.35,y+.15,.4,.22,11,GOLD,True); tb(s,cap,2.85,y+.15,4,.25,14,WHITE,True)
slide_transition(s); footer(s,18)
# end
s=slide(); bg(s); pic(s,by['loess'][1] if len(by['loess'])>1 else by['loess'][0],0,0,W,H); rect(s,0,0,W,H,BLACK,.58); tb(s,'邢堡村值得被看见',1,2.45,9,.8,44,WHITE,True,font='SimSun'); gold(s,1.02,3.45,1.15); tb(s,'让黄土塬、古院落、麻花香气和板枣甜意，成为山西乡村文旅的新叙事。',1.02,3.85,8.5,.6,20,CREAM); tb(s,'Thank you',9.25,6.75,3,.35,14,GOLD,True,PP_ALIGN.RIGHT); slide_transition(s); footer(s,19)
prs.save(OUT)
print(f'Saved {OUT}: {len(prs.slides)} slides, items {len(items)}')
